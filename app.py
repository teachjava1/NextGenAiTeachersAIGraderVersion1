# -*- coding: utf-8 -*-
import os
import io
import sqlite3
import hashlib
import stripe
from datetime import date, datetime

from flask import (
    Flask, request, jsonify, send_from_directory,
    session
)
from openai import OpenAI

from docx import Document as DocxDocument
from PyPDF2 import PdfReader
import openpyxl
import re

# ------------ Config ------------

DATABASE_PATH = "nextgen_ai_teachers_aigrader.db"
FREE_DAILY_LIMIT = 5  # free plan limit

client = OpenAI()  # uses OPENAI_API_KEY from environment

stripe.api_key = os.environ.get("STRIPE_SECRET_KEY")

app = Flask(__name__, static_folder="static", static_url_path="")
app.secret_key = os.environ.get("FLASK_SECRET_KEY", "dev-secret-change-me")  # replace in production

# ------------ Default Prompt ------------

DEFAULT_PROMPT_TEMPLATE = """You are a rubric-aware grading assistant.

You will receive two things:
1. A rubric (can be a list, table, or plain text)
2. A student submission

Your job is to produce clear, plain-text feedback with NO Markdown formatting.

IMPORTANT RULES:
- Do NOT use ** or any markdown formatting.
- Do NOT use bullets like *, -, +.
- Do NOT use horizontal rules such as ---.
- Do NOT include backticks or code fences.
- Use plain text only: letters, numbers, parentheses, colons, and normal sentences.

STRUCTURE THE OUTPUT EXACTLY LIKE THIS:

1. <Category Name> (<MaxPoints> points)
Score: <EarnedPoints>/<MaxPoints>
Justification: <short explanation in 1–3 sentences>

(blank line)

2. <Next Category> (<MaxPoints> points)
Score: <EarnedPoints>/<MaxPoints>
Justification: <short explanation>

(blank line between every category)

At the end, ALWAYS include this final section:

Teacher Comment Summary:
<1–3 sentences summarizing strengths and improvements in plain text, ready to paste into Canvas or Google Classroom>

Rubric:
{rubric}

Student Submission:
{student}"""

# ------------ Simple SHA256 Password Hashing ------------

def hash_password(password: str) -> str:
    """
    Hash a password using SHA256.
    NOTE: This is a simple hash, not salted. Fine for a small demo / local app.
    """
    return hashlib.sha256(password.encode("utf-8")).hexdigest()


def verify_password(password: str, stored_hash: str) -> bool:
    """
    Compare a plain-text password to a SHA256 hash.
    """
    return hash_password(password) == stored_hash

# ------------ DB Helpers ------------

def get_db():
    conn = sqlite3.connect(DATABASE_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = get_db()
    cur = conn.cursor()

    cur.execute("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            email TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            plan TEXT NOT NULL DEFAULT 'free', -- 'free' or 'pro'
            uses_today INTEGER NOT NULL DEFAULT 0,
            last_use_date TEXT
        );
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS activation_codes (
            code TEXT PRIMARY KEY,
            plan TEXT NOT NULL,             -- e.g. 'pro'
            redeemed_by INTEGER,            -- user id
            redeemed_at TEXT,               -- ISO timestamp
            FOREIGN KEY(redeemed_by) REFERENCES users(id)
        );
    """)

    conn.commit()
    conn.close()


# Call at startup
init_db()

# ------------ File Text Extraction ------------

def extract_text_from_file(file_storage):
    filename = file_storage.filename or ""
    ext = os.path.splitext(filename)[1].lower()
    raw = file_storage.read()
    file_storage.seek(0)

    # Plain text / code / CSV
    if ext in [".txt", ".cpp", ".java", ".py", ".md", ".xml", ".html", ".json", ".csv"]:
        try:
            return raw.decode("utf-8", errors="ignore")
        except Exception:
            return "[Error decoding text file.]"

    # DOCX
    if ext == ".docx":
        try:
            buffer = io.BytesIO(raw)
            doc = DocxDocument(buffer)
            return "\n".join(p.text for p in doc.paragraphs)
        except Exception:
            return "[Error reading DOCX file.]"

    # PDF
    if ext == ".pdf":
        try:
            buffer = io.BytesIO(raw)
            reader = PdfReader(buffer)
            pages = []
            for page in reader.pages:
                text = page.extract_text() or ""
                pages.append(text)
            return "\n".join(pages)
        except Exception:
            return "[Error reading PDF file.]"

    # XLSX
    if ext == ".xlsx":
        try:
            buffer = io.BytesIO(raw)
            wb = openpyxl.load_workbook(buffer, data_only=True)
            sheet = wb.active
            lines = []
            for row in sheet.iter_rows(values_only=True):
                line = "\t".join("" if cell is None else str(cell) for cell in row)
                lines.append(line)
            return "\n".join(lines)
        except Exception:
            return "[Error reading XLSX file.]"

    return f"[Unsupported or unknown file type: {ext}]"

# ------------ AI Call ------------

def call_model(full_prompt: str) -> str:
    """Call OpenAI and return plain text."
    Uses chat.completions for maximum compatibility.
    """
    response = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[{"role": "user", "content": full_prompt}],
        temperature=0.0,
    )
    return response.choices[0].message.content

# ------------ Auth Helpers ------------

def get_current_user():
    user_id = session.get("user_id")
    if not user_id:
        return None
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT * FROM users WHERE id = ?", (user_id,))
    row = cur.fetchone()
    conn.close()
    return row


def update_user_usage(user_id):
    today_str = date.today().isoformat()
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT uses_today, last_use_date, plan FROM users WHERE id = ?", (user_id,))
    row = cur.fetchone()
    if not row:
        conn.close()
        return {"error": "User not found"}

    uses_today = row["uses_today"] or 0
    last_use_date = row["last_use_date"]
    plan = row["plan"]

    if last_use_date != today_str:
        uses_today = 0

    if plan == "free" and uses_today >= FREE_DAILY_LIMIT:
        conn.close()
        return {"error": "Free plan daily limit reached. Upgrade to Pro for unlimited grading."}

    uses_today += 1
    cur.execute("""
        UPDATE users SET uses_today = ?, last_use_date = ? WHERE id = ?
    """, (uses_today, today_str, user_id))
    conn.commit()
    conn.close()

    return {
        "plan": plan,
        "uses_today": uses_today,
        "limit": FREE_DAILY_LIMIT
    }

# ------------ Routes ------------

@app.route("/")
def index():
    return send_from_directory(app.static_folder, "index.html")

# --- Auth APIs ---

@app.post("/api/register")
def register():
    data = request.get_json() or {}
    email = (data.get("email") or "").strip().lower()
    password = (data.get("password") or "").strip()

    if not email or not password:
        return jsonify({"error": "Email and password are required."}), 400

    pw_hash = hash_password(password)

    conn = get_db()
    cur = conn.cursor()
    try:
        cur.execute("""
            INSERT INTO users (email, password_hash, plan)
            VALUES (?, ?, 'free')
        """, (email, pw_hash))
        conn.commit()
    except sqlite3.IntegrityError:
        conn.close()
        return jsonify({"error": "Email already registered."}), 400

    cur.execute("SELECT id FROM users WHERE email = ?", (email,))
    row = cur.fetchone()
    conn.close()

    session["user_id"] = row["id"]
    return jsonify({"message": "Registered successfully.", "plan": "free"})


@app.post("/api/login")
def login():
    data = request.get_json() or {}
    email = (data.get("email") or "").strip().lower()
    password = (data.get("password") or "").strip()

    if not email or not password:
        return jsonify({"error": "Email and password are required."}), 400

    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT * FROM users WHERE email = ?", (email,))
    row = cur.fetchone()
    conn.close()

    if not row or not verify_password(password, row["password_hash"]):
        return jsonify({"error": "Invalid email or password."}), 400

    session["user_id"] = row["id"]
    return jsonify({
        "message": "Logged in successfully.",
        "plan": row["plan"]
    })


@app.post("/api/logout")
def logout():
    session.pop("user_id", None)
    return jsonify({"message": "Logged out."})


@app.get("/api/me")
def me():
    user = get_current_user()
    if not user:
        return jsonify({"loggedIn": False})
    return jsonify({
        "loggedIn": True,
        "email": user["email"],
        "plan": user["plan"],
        "uses_today": user["uses_today"],
        "limit": FREE_DAILY_LIMIT
    })

# --- Activation Code Redeem ---

@app.post("/api/redeem")
def redeem():
    user = get_current_user()
    if not user:
        return jsonify({"error": "Login required."}), 401

    data = request.get_json() or {}
    code = (data.get("code") or "").strip()

    if not code:
        return jsonify({"error": "Activation code is required."}), 400

    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT * FROM activation_codes WHERE code = ?", (code,))
    code_row = cur.fetchone()

    if not code_row:
        conn.close()
        return jsonify({"error": "Invalid activation code."}), 400

    if code_row["redeemed_by"] is not None:
        conn.close()
        return jsonify({"error": "Code has already been redeemed."}), 400

    # Upgrade user plan
    cur.execute("UPDATE users SET plan = ? WHERE id = ?", (code_row["plan"], user["id"]))
    cur.execute("""
        UPDATE activation_codes
        SET redeemed_by = ?, redeemed_at = ?
        WHERE code = ?
    """, (user["id"], datetime.utcnow().isoformat(), code))
    conn.commit()
    conn.close()

    return jsonify({"message": f"Code redeemed! Your plan is now {code_row['plan']}."})


@app.post("/api/create-checkout-session")
def create_checkout_session():
    """
    Start a Stripe Checkout session for upgrading to Pro.
    Uses STRIPE_PRICE_ID from environment for the subscription price.
    """
    user = get_current_user()
    if not user:
        return jsonify({"error": "Login required."}), 401

    price_id = os.environ.get("STRIPE_PRICE_ID")
    if not price_id:
        return jsonify({"error": "Server not configured with STRIPE_PRICE_ID."}), 500

    try:
        success_url = request.host_url.rstrip("/") + "/success"
        cancel_url = request.host_url.rstrip("/") + "/cancel"

        session_obj = stripe.checkout.Session.create(
            payment_method_types=["card"],
            mode="subscription",
            line_items=[{"price": price_id, "quantity": 1}],
            success_url=success_url,
            cancel_url=cancel_url,
            client_reference_id=str(user["id"]),
        )

        return jsonify({
            "sessionId": session_obj.id,
            "publicKey": os.environ.get("STRIPE_PUBLIC_KEY", "")
        })
    except Exception as e:
        return jsonify({"error": f"Stripe error: {e}"}), 500


@app.post("/api/stripe-webhook")
def stripe_webhook():
    """
    Handle Stripe webhook events.
    Used to upgrade users to Pro after successful checkout.
    """
    payload = request.data
    sig = request.headers.get("Stripe-Signature")
    webhook_secret = os.environ.get("STRIPE_WEBHOOK_SECRET")

    try:
        if webhook_secret:
            event = stripe.Webhook.construct_event(
                payload, sig, webhook_secret
            )
        else:
            event = json.loads(payload.decode("utf-8"))
    except Exception as e:
        return jsonify({"error": f"Webhook error: {e}"}), 400

    if event.get("type") == "checkout.session.completed":
        session_obj = event["data"]["object"]
        user_id = session_obj.get("client_reference_id")
        if user_id:
            conn = get_db()
            cur = conn.cursor()
            cur.execute("UPDATE users SET plan = 'pro' WHERE id = ?", (user_id,))
            conn.commit()
            conn.close()

    return jsonify({"status": "ok"})


@app.get("/success")
def success():
    return "Payment successful! Your account is now Pro. You can close this tab and return to the app."


@app.get("/cancel")
def cancel():
    return "Payment canceled. You can close this tab and return to the app."


# --- Grading API ---

@app.post("/api/grade")
def grade():
    user = get_current_user()
    if not user:
        return jsonify({"error": "Login required."}), 401

    usage_info = update_user_usage(user["id"])
    if "error" in usage_info:
        return jsonify(usage_info), 403

    rubric_text = (request.form.get("rubricText") or "").strip()
    student_text = (request.form.get("studentText") or "").strip()
    prompt_template = (request.form.get("promptTemplate") or "").strip() or DEFAULT_PROMPT_TEMPLATE

    rubric_file = request.files.get("rubricFile")
    student_file = request.files.get("studentFile")

    if rubric_file and rubric_file.filename:
        rubric_text = extract_text_from_file(rubric_file)

    if student_file and student_file.filename:
        student_text = extract_text_from_file(student_file)

    if not rubric_text or not student_text:
        return jsonify({"error": "Rubric and student submission are required (via file or text)."}), 400

    if "{rubric}" not in prompt_template or "{student}" not in prompt_template:
        return jsonify({"error": "Prompt template must contain {rubric} and {student}."}), 400

    full_prompt = (
        prompt_template
        .replace("{rubric}", rubric_text)
        .replace("{student}", student_text)
    )

    try:
        result_text = call_model(full_prompt)
    except Exception as e:
        return jsonify({"error": f"OpenAI error: {e}"}), 500

    # Clean up any leftover markdown-like symbols so output is plain text
    clean_result = result_text
    clean_result = clean_result.replace("**", "")
    clean_result = clean_result.replace("---", "")
    import re as _re_local
    clean_result = _re_local.sub(r"^[*\-\+]+\s*", "", clean_result, flags=_re_local.MULTILINE)
    clean_result = clean_result.strip()

    return jsonify({
        "result": clean_result,
        "plan": usage_info["plan"],
        "uses_today": usage_info["uses_today"],
        "limit": usage_info["limit"]
    })


if __name__ == "__main__":
    app.run(debug=True)
